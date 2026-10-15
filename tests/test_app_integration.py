'''
Integration tests for app.py (the local control panel) via Flask's test client.
These exercise real request/response behaviour: schema exposure, config save
coercion + round-trip, unknown-key rejection, and the applied-jobs history
CSV -> JSON mapping and mark-applied flow.

All tests isolate their writes to a tmp_path, so the user's real user_config.json
and "all excels/" folder are never touched.

License: MIT  (https://opensource.org/license/mit)
'''

import csv
import io
import json
import os


# --------------------------------- schema -----------------------------------
def test_schema_endpoint_returns_nonempty_list(client):
    resp = client.get("/api/schema")
    assert resp.status_code == 200
    data = resp.get_json()
    assert isinstance(data, list) and len(data) > 0
    assert "section" in data[0] and "fields" in data[0]


def test_schema_includes_ollama_gemma_setup(client):
    sections = client.get("/api/schema").get_json()
    account = next(section for section in sections if section["section"] == "Account")
    fields = {field["key"]: field for field in account["fields"]}

    assert "openai" in fields["ai_provider"]["options"]
    suggested = fields["llm_model"]["models_by_provider"]["openai"]
    assert {"gemma4:e2b", "gemma4:e4b", "gemma4:12b", "gemma4:26b"}.issubset(suggested)
    assert {"gemma3:4b", "gemma3:12b", "gemma3:27b", "qwen3.5:4b"}.issubset(suggested)
    assert "localhost:11434/v1/" in fields["llm_api_url"]["help"]


def test_playwright_recorder_accepts_only_local_draft_urls():
    from modules.playwright_local import is_http_url, is_local_draft_url

    assert is_local_draft_url("http://127.0.0.1:51048/draft-application")
    assert is_local_draft_url("http://localhost:51048/draft-application/")
    assert not is_local_draft_url("https://localhost:51048/draft-application")
    assert not is_local_draft_url("http://localhost:51048/")
    assert not is_local_draft_url("http://localhost:51048/draft-application?next=external")
    assert not is_local_draft_url("https://www.linkedin.com/jobs/")
    assert is_http_url("https://www.linkedin.com/jobs/view/123")
    assert not is_http_url("javascript:alert(1)")
    assert not is_http_url("https://user:password@example.com/job")


def test_strict_json_profile_import_returns_review_preview_without_saving(client, tmp_path, monkeypatch):
    import app
    cfg_path = str(tmp_path / "user_config.json")
    monkeypatch.setattr(app, "USER_CONFIG_PATH", cfg_path)
    payload = {
        "personals": {"first_name": "Ada", "last_name": "Lovelace"},
        "questions": {"years_of_experience": "5", "desired_salary": 120000},
    }
    response = client.post(
        "/api/import-profile",
        data={"file": (io.BytesIO(json.dumps(payload).encode()), "profile.json")},
        content_type="multipart/form-data",
    )
    assert response.status_code == 200
    result = response.get_json()
    assert result["source"] == "json"
    values = {(field["module"], field["key"]): field["value"] for field in result["fields"]}
    assert values[("personals", "first_name")] == "Ada"
    assert values[("questions", "desired_salary")] == 120000
    assert not os.path.exists(cfg_path)


def test_strict_json_profile_import_rejects_unknown_fields(client):
    payload = {"personals": {"first_name": "Ada", "unknown": "value"}, "questions": {}}
    response = client.post(
        "/api/import-profile",
        data={"file": (io.BytesIO(json.dumps(payload).encode()), "profile.json")},
        content_type="multipart/form-data",
    )
    assert response.status_code == 400
    assert "Unknown personals" in response.get_json()["error"]


def test_profile_import_extracts_docx_text():
    from docx import Document
    import app

    doc = Document()
    doc.add_paragraph("Ada Lovelace")
    doc.add_paragraph("Phone: +61 400 000 000")
    buffer = io.BytesIO()
    doc.save(buffer)
    extracted = app._extract_profile_document("candidate.docx", buffer.getvalue())
    assert "Ada Lovelace" in extracted
    assert "+61 400 000 000" in extracted


def test_session_application_tracker_is_volatile_and_supports_human_status(client):
    import app
    with app._SESSION_APPLICATIONS_LOCK:
        app._SESSION_APPLICATIONS.clear()
    created = client.post("/api/session-applications", json={
        "job_title": "Data Analyst", "company": "Example Co", "status": "needs_user_input",
        "notes": "Confirm notice period", "job_url": "https://example.com/jobs/1",
    })
    assert created.status_code == 201
    entry = created.get_json()
    assert client.get("/api/session-applications").get_json()[0]["status"] == "needs_user_input"

    updated = client.put("/api/session-applications/" + entry["id"], json={
        "job_title": "Data Analyst", "company": "Example Co", "status": "applied",
        "notes": "Submitted manually", "job_url": "https://example.com/jobs/1",
    })
    assert updated.status_code == 200
    assert updated.get_json()["status"] == "applied"
    assert client.delete("/api/session-applications/" + entry["id"]).status_code == 200
    assert client.get("/api/session-applications").get_json() == []


def test_workflow_page_renders_and_browser_assist_rejects_invalid_url(client):
    assert client.get("/workflow").status_code == 200
    response = client.post("/api/browser-assist", json={"url": "javascript:alert(1)"})
    assert response.status_code == 400


# ------------------------------- config save/get ----------------------------
def test_config_save_coerces_and_roundtrips(client, tmp_path, monkeypatch):
    import app
    import config._overrides as overrides
    cfg_path = str(tmp_path / "user_config.json")
    monkeypatch.setattr(app, "USER_CONFIG_PATH", cfg_path)
    monkeypatch.setattr(overrides, "USER_CONFIG_PATH", cfg_path)

    # "true" (string) must be coerced to a real bool for the use_AI field.
    resp = client.post("/api/config", json={"secrets": {"use_AI": "true"}})
    assert resp.status_code == 200

    with open(cfg_path, encoding="utf-8") as f:
        saved = json.load(f)
    assert saved["secrets"]["use_AI"] is True

    # GET reflects the saved value.
    got = client.get("/api/config").get_json()
    assert got["secrets"]["use_AI"] is True


def test_config_save_rejects_unknown_key(client, tmp_path, monkeypatch):
    import app
    import config._overrides as overrides
    cfg_path = str(tmp_path / "user_config.json")
    monkeypatch.setattr(app, "USER_CONFIG_PATH", cfg_path)
    monkeypatch.setattr(overrides, "USER_CONFIG_PATH", cfg_path)

    resp = client.post("/api/config", json={"secrets": {"definitely_not_a_field": 1}})
    assert resp.status_code == 400
    assert not os.path.exists(cfg_path)  # nothing written on rejection


def test_config_save_rejects_unknown_section(client, tmp_path, monkeypatch):
    import app
    import config._overrides as overrides
    cfg_path = str(tmp_path / "user_config.json")
    monkeypatch.setattr(app, "USER_CONFIG_PATH", cfg_path)
    monkeypatch.setattr(overrides, "USER_CONFIG_PATH", cfg_path)

    resp = client.post("/api/config", json={"not_a_section": {"x": 1}})
    assert resp.status_code == 400


# ------------------------------ applied-jobs CSV ----------------------------
_CSV_COLUMNS = ['Job ID', 'Title', 'Company', 'HR Name', 'HR Link',
                'Job Link', 'External Job link', 'Date Applied']


def _write_history_csv(folder):
    path = os.path.join(folder, "all_applied_applications_history.csv")
    with open(path, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=_CSV_COLUMNS)
        writer.writeheader()
        writer.writerow({
            'Job ID': 'J1', 'Title': 'Engineer', 'Company': 'Acme',
            'HR Name': 'Unknown', 'HR Link': '', 'Job Link': 'http://x',
            'External Job link': 'http://ext', 'Date Applied': 'Pending',
        })
    return path


def test_applied_jobs_get_maps_columns_to_json_keys(client, tmp_path, monkeypatch):
    import app
    monkeypatch.setattr(app, "PATH", str(tmp_path) + os.sep)
    _write_history_csv(str(tmp_path))

    resp = client.get("/applied-jobs")
    assert resp.status_code == 200
    row = resp.get_json()[0]
    assert row["Job_ID"] == "J1"
    assert row["Title"] == "Engineer"
    assert row["External_Job_link"] == "http://ext"
    assert row["Date_Applied"] == "Pending"


def test_applied_jobs_mark_applied_updates_date(client, tmp_path, monkeypatch):
    import app
    monkeypatch.setattr(app, "PATH", str(tmp_path) + os.sep)
    _write_history_csv(str(tmp_path))

    resp = client.put("/applied-jobs/J1")
    assert resp.status_code == 200

    row = client.get("/applied-jobs").get_json()[0]
    assert row["Date_Applied"] != "Pending"


def test_applied_jobs_missing_file_returns_404(client, tmp_path, monkeypatch):
    import app
    monkeypatch.setattr(app, "PATH", str(tmp_path) + os.sep)  # empty dir, no CSV
    assert client.get("/applied-jobs").status_code == 404


def test_applied_jobs_mark_unknown_id_returns_404(client, tmp_path, monkeypatch):
    import app
    monkeypatch.setattr(app, "PATH", str(tmp_path) + os.sep)
    _write_history_csv(str(tmp_path))
    assert client.put("/applied-jobs/does-not-exist").status_code == 404


# ------------------------------- bot status ---------------------------------
def test_status_reports_not_running(client):
    resp = client.get("/api/status")
    assert resp.status_code == 200
    assert resp.get_json()["running"] is False


# ------------------------------- pages render -------------------------------
def test_control_panel_and_history_pages_render(client):
    assert client.get("/").status_code == 200
    assert client.get("/history").status_code == 200
