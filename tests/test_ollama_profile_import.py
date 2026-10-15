"""Optional real-Ollama smoke test for structured résumé/profile extraction."""

import io
import os

import pytest


@pytest.mark.live
@pytest.mark.skipif(not os.getenv("OLLAMA_TEST_MODEL"), reason="set OLLAMA_TEST_MODEL to run the local Ollama smoke test")
def test_live_ollama_extracts_profile_from_docx(client, monkeypatch):
    from docx import Document
    import app
    import config.secrets as cfg

    model_name = os.environ["OLLAMA_TEST_MODEL"]
    monkeypatch.setattr(app.importlib, "reload", lambda module: module)
    monkeypatch.setattr(cfg, "use_AI", True)
    monkeypatch.setattr(cfg, "ai_provider", "openai")
    monkeypatch.setattr(cfg, "llm_model", model_name)
    monkeypatch.setattr(cfg, "llm_api_key", "ollama")
    monkeypatch.setattr(cfg, "llm_api_url", "http://127.0.0.1:11434/v1/")

    resume = Document()
    resume.add_paragraph("Ada Lovelace")
    resume.add_paragraph("Phone: +61 400 000 000")
    resume.add_paragraph("5 years of professional experience.")
    content = io.BytesIO()
    resume.save(content)

    response = client.post(
        "/api/import-profile",
        data={
            "file": (io.BytesIO(content.getvalue()), "candidate.docx"),
            "provider_consent": "true",
        },
        content_type="multipart/form-data",
    )

    assert response.status_code == 200, response.get_json()
    result = response.get_json()
    values = {(field["module"], field["key"]): field["value"] for field in result["fields"]}
    assert values[("personals", "first_name")].lower() == "ada", result
    assert values[("personals", "last_name")].lower() == "lovelace"
    assert str(values[("questions", "years_of_experience")]).strip().lower() in {"5", "five"}
